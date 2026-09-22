#!/usr/bin/env python3
"""Generate PERSONAS.docx из PERSONAS.md по стандарту Алексея Краюшкина."""
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


# ─────────────────────────────────────────────────────────────────────────
# Парсер PERSONAS.md
# ─────────────────────────────────────────────────────────────────────────

def parse_md(path: Path):
    text = path.read_text(encoding="utf-8")
    lines = text.split("\n")

    blocks = []
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()

        if not ln.strip():
            i += 1
            continue

        if re.match(r"^#\s+", ln) and not re.match(r"^##\s+", ln):
            i += 1
            continue

        m = re.match(r"^(#{2,4})\s+(.*)$", ln)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            blocks.append((f"h{level}", title))
            i += 1
            continue

        if ln.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|", lines[i + 1].strip()):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            rows = []
            for row in tbl:
                if re.match(r"^\|[\s\-:|]+\|$", row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                rows.append(cells)
            blocks.append(("table", rows))
            continue

        if ln.lstrip().startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote_lines.append(lines[i].lstrip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(quote_lines)))
            continue

        if re.match(r"^\d+\.\s+", ln):
            items = []
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[i]).rstrip())
                i += 1
            blocks.append(("ol", items))
            continue

        if re.match(r"^\s*-\s+", ln) or re.match(r"^\s*\*\s+", ln):
            items = []
            while i < len(lines) and (re.match(r"^\s*-\s+", lines[i]) or re.match(r"^\s*\*\s+", lines[i])):
                raw = lines[i]
                indent = len(raw) - len(raw.lstrip())
                txt = re.sub(r"^\s*[-\*]\s+", "", raw).rstrip()
                items.append((indent, txt))
                i += 1
            blocks.append(("ul", items))
            continue

        para_lines = []
        while i < len(lines):
            cur = lines[i].rstrip()
            if not cur:
                break
            if (cur.startswith("#") or cur.lstrip().startswith(">")
                    or cur.strip().startswith("|") or re.match(r"^\s*[-*]\s+", cur)
                    or re.match(r"^\s*\d+\.\s+", cur)):
                break
            para_lines.append(cur)
            i += 1
        if para_lines:
            blocks.append(("p", " ".join(para_lines)))

    return blocks


# ─────────────────────────────────────────────────────────────────────────
# Очистка текста
# ─────────────────────────────────────────────────────────────────────────

def clean_text(s: str) -> str:
    s = s.replace("—", "-").replace("–", "-")
    s = s.replace("\xa0", " ")
    s = re.sub(r" {2,}", " ", s)
    s = re.sub(r"\n{2,}", "\n", s)
    return s.strip()


# ─────────────────────────────────────────────────────────────────────────
# Стили (стандарт Краюшкина)
# ─────────────────────────────────────────────────────────────────────────

def setup_normal_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    for old in rPr.findall(qn("w:lang")):
        rPr.remove(old)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")
    rPr.append(lang)
    pf = normal.paragraph_format
    pf.first_line_indent = Mm(12.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def setup_caption_style(doc):
    styles = doc.styles
    try:
        cap = styles["Подпись"]
    except KeyError:
        cap = styles.add_style("Подпись", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = styles["Normal"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(12)
    cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    pf = cap.paragraph_format
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_with_next = True


def set_page_margins(doc):
    for section in doc.sections:
        section.left_margin = Mm(20)
        section.right_margin = Mm(10)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)


def setup_headings(doc):
    sizes = {1: 16, 2: 14, 3: 12}
    for lvl, sz in sizes.items():
        try:
            s = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "Times New Roman")
        rFonts.set(qn("w:cs"), "Times New Roman")
        pf = s.paragraph_format
        pf.first_line_indent = Mm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


# ─────────────────────────────────────────────────────────────────────────
# Добавление элементов
# ─────────────────────────────────────────────────────────────────────────

def add_document_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(clean_text(text))
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p


def add_list_bullet(doc, items):
    for indent, txt in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(7.5 * (1 + indent // 2))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_number(doc, items):
    for txt in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Mm(10)
    p.paragraph_format.right_indent = Mm(10)
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    # рамка сверху и снизу для выделения цитаты
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "bottom"):
        bdr = OxmlElement(f"w:{side}")
        bdr.set(qn("w:val"), "single")
        bdr.set(qn("w:sz"), "4")
        bdr.set(qn("w:space"), "4")
        bdr.set(qn("w:color"), "888888")
        pBdr.append(bdr)
    pPr.append(pBdr)


def add_seq_field(paragraph, name="Таблица"):
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {name} \\* ARABIC "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def add_table_caption_above(doc, caption_text):
    p = doc.add_paragraph(style="Подпись")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run1 = p.add_run(f"{clean_text(caption_text)} приведено в таблице (Таблица ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Таблица")
    run2 = p.add_run(").")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_table(doc, rows):
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Light Grid Accent 1"
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean_text(cell_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return table


# ─────────────────────────────────────────────────────────────────────────
# Пост-обработка DOCX
# ─────────────────────────────────────────────────────────────────────────

def post_process_docx(docx_path: Path) -> bool:
    tmp = docx_path.with_suffix(".tmp.docx")
    seen_n = False
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/numbering.xml":
                    if seen_n:
                        continue
                    seen_n = True
                zout.writestr(item.filename, zin.read(item.filename))
    shutil.move(tmp, docx_path)
    with zipfile.ZipFile(docx_path, "r") as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(name))
                except ET.ParseError as e:
                    print(f"INVALID: {name}: {e}")
                    return False
    return True


# ─────────────────────────────────────────────────────────────────────────
# Сборка документа
# ─────────────────────────────────────────────────────────────────────────

def strip_prefix(text):
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", text).strip()


def add_paragraph_with_emojis(doc, text):
    """Добавляет параграф, сохраняя эмодзи (которые могут быть в тексте)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p


def build_docx(md_path: Path, out_path: Path, project_name: str):
    blocks = parse_md(md_path)
    doc = Document()
    setup_normal_style(doc)
    setup_headings(doc)
    setup_caption_style(doc)
    set_page_margins(doc)

    doc.core_properties.author = "Алексей Краюшкин"
    doc.core_properties.title = "User Personas HTML_MeetingProtokol"

    add_document_title(doc, "User Personas HTML_MeetingProtokol")

    table_counter = 0
    for kind, payload in blocks:
        if kind == "h2":
            add_heading(doc, strip_prefix(payload), 1)
        elif kind == "h3":
            add_heading(doc, strip_prefix(payload), 2)
        elif kind == "h4":
            add_heading(doc, strip_prefix(payload), 3)
        elif kind == "p":
            add_paragraph_with_emojis(doc, payload)
        elif kind == "ul":
            add_list_bullet(doc, payload)
        elif kind == "ol":
            add_list_number(doc, payload)
        elif kind == "quote":
            add_quote(doc, payload)
        elif kind == "table":
            table_counter += 1
            table_names = {
                1: "Метаинформация о документе",
                2: "Сводная карта персон проекта",
                3: "Демографический профиль персоны",
                4: "Психографический профиль",
                5: "Техническая грамотность",
                6: "Контекст использования",
                7: "Боли и проблемы персоны",
                8: "Цели персоны",
                9: "Типичные сценарии использования",
                10: "Empathy Map персоны",
                11: "Технические ограничения",
                12: "Каналы коммуникации",
                13: "Связь персоны с фичами продукта",
                14: "Матрица связей фич и персон",
            }
            caption_text = table_names.get(table_counter, f"Сводные данные {table_counter}")
            add_table_caption_above(doc, caption_text)
            add_table(doc, payload)

    add_page_numbers(doc)
    doc.save(out_path)
    ok = post_process_docx(out_path)
    return ok


def main():
    project_root = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
    md_path = project_root / "artifacts" / "02-personas" / "PERSONAS.md"
    if not md_path.exists():
        print(f"Not found: {md_path}", file=sys.stderr)
        sys.exit(1)
    out_path = md_path.parent / "PERSONAS.docx"
    project_name = project_root.name
    ok = build_docx(md_path, out_path, project_name)
    if ok:
        print(f"OK: {out_path} ({out_path.stat().st_size} bytes)")
    else:
        print("FAIL: XML-валидность нарушена", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()