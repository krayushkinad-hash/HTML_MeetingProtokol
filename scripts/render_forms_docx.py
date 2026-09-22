#!/usr/bin/env python3
"""Generate FORMS_SPEC.docx — сводный документ по всем формам HTML_MeetingProtokol."""
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


# ─────────────────────────────────────────────────────────────────────────
# Очистка текста и стили (Краюшкин)
# ─────────────────────────────────────────────────────────────────────────

def clean_text(s):
    s = s.replace("—", "-").replace("–", "-")
    s = s.replace("\xa0", " ")
    s = re.sub(r" {2,}", " ", s)
    return s.strip()


def setup_normal_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    for old in rPr.findall(qn("w:lang")):
        rPr.remove(old)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")
    rPr.append(lang)
    pf = normal.paragraph_format
    pf.first_line_indent = Mm(12.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def setup_caption_style(doc):
    styles = doc.styles
    try:
        cap = styles["Подпись"]
    except KeyError:
        cap = styles.add_style("Подпись", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = styles["Normal"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(12)
    cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    pf = cap.paragraph_format
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_with_next = True


def set_page_margins(doc):
    for section in doc.sections:
        section.left_margin = Mm(20)
        section.right_margin = Mm(10)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)


def setup_headings(doc):
    sizes = {1: 16, 2: 14, 3: 12}
    for lvl, sz in sizes.items():
        try:
            s = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "Times New Roman")
        rFonts.set(qn("w:cs"), "Times New Roman")
        pf = s.paragraph_format
        pf.first_line_indent = Mm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


# ─────────────────────────────────────────────────────────────────────────
# Добавление элементов
# ─────────────────────────────────────────────────────────────────────────

def add_document_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(clean_text(text))


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_bullet(doc, items):
    for indent, txt in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(7.5 * (1 + indent // 2))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_number(doc, items):
    for txt in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_seq_field(paragraph, name):
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {name} \\* ARABIC "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def add_table_caption_above(doc, caption_text):
    p = doc.add_paragraph(style="Подпись")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run1 = p.add_run(f"{clean_text(caption_text)} приведено в таблице (Таблица ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Таблица")
    run2 = p.add_run(").")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_table(doc, rows):
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Light Grid Accent 1"
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean_text(cell_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return table


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def add_code_block(doc, code_text):
    """Consolas 9 пт — для блоков кода/JSON внутри документа."""
    for line in code_text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(12.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


# ─────────────────────────────────────────────────────────────────────────
# Парсер FORM_*.md
# ─────────────────────────────────────────────────────────────────────────

def parse_form_md(md_path):
    """Парсит FORM_<Code>.md, возвращает список секций (h2, table, bullet, code, p)."""
    text = md_path.read_text(encoding="utf-8")
    blocks = []
    i = 0
    lines = text.split("\n")

    in_code = False
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip()

        # Code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                blocks.append(("code", "\n".join(code_lines)))
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # H1 — пропускаем (заголовок документа свой)
        if re.match(r"^#\s+", line) and not re.match(r"^##\s+", line):
            i += 1
            continue

        # H2 / H3
        m = re.match(r"^(#{2,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            blocks.append((f"h{level}", title))
            i += 1
            continue

        # Таблица
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|", lines[i + 1].strip()):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            rows = []
            for row in tbl:
                if re.match(r"^\|[\s\-:|]+\|$", row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                rows.append(cells)
            blocks.append(("table", rows))
            continue

        # Буллет
        if re.match(r"^\s*-\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*-\s+", lines[i]):
                raw = lines[i]
                indent = len(raw) - len(raw.lstrip())
                txt = re.sub(r"^\s*-\s+", "", raw).rstrip()
                items.append((indent, txt))
                i += 1
            blocks.append(("ul", items))
            continue

        # Параграф
        if line.strip():
            blocks.append(("p", line))
        i += 1

    return blocks


# ─────────────────────────────────────────────────────────────────────────
# Пост-обработка DOCX
# ─────────────────────────────────────────────────────────────────────────

def post_process_docx(docx_path):
    tmp = docx_path.with_suffix(".tmp.docx")
    seen_n = False
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/numbering.xml":
                    if seen_n:
                        continue
                    seen_n = True
                zout.writestr(item.filename, zin.read(item.filename))
    shutil.move(tmp, docx_path)
    with zipfile.ZipFile(docx_path, "r") as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(name))
                except ET.ParseError as e:
                    print(f"INVALID: {name}: {e}")
                    return False
    return True


# ─────────────────────────────────────────────────────────────────────────
# Сборка документа
# ─────────────────────────────────────────────────────────────────────────

def strip_prefix(text):
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", text).strip()


def build_docx(forms_dir, out_path):
    form_files = sorted(forms_dir.glob("FORM_*.md"))
    print(f"Найдено {len(form_files)} файлов форм")

    doc = Document()
    setup_normal_style(doc)
    setup_headings(doc)
    setup_caption_style(doc)
    set_page_margins(doc)

    doc.core_properties.author = "Алексей Краюшкин"
    doc.core_properties.title = "Спецификации форм HTML_MeetingProtokol"

    # ── Титульный лист ──
    add_document_title(doc, "Спецификации форм HTML_MeetingProtokol")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("Сводный документ по всем формам приложения")
    run.font.name = "Times New Roman"; run.font.size = Pt(14); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("Версия 1.0 · 2026-09-14")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_page_break(doc)

    # ── §1. Метаинформация ──
    add_heading(doc, "1 Метаинформация", 1)
    add_table_caption_above(doc, "Метаинформация документа")
    add_table(doc, [
        ["Поле", "Значение"],
        ["Проект", "HTML_MeetingProtokol"],
        ["Документ", "Спецификации форм"],
        ["Версия", "1.0"],
        ["Дата", "2026-09-14"],
        ["Формат", "DOCX по стандарту Алексея Краюшкина"],
        ["Всего форм", str(len(form_files))],
        ["Автор", "Алексей Краюшкин"],
    ])

    # ── §2. Реестр форм ──
    add_heading(doc, "2 Реестр форм", 1)
    add_table_caption_above(doc, "Сводный реестр форм проекта")
    registry_rows = [["ID формы", "Файл", "Название", "Размер (КБ)"]]
    for f in form_files:
        form_id = "F-HMP-" + f.stem.replace("FORM_F-HMP-", "").replace("FORM_", "")
        # Парсим ID из первой строки таблицы метаинформации
        text = f.read_text(encoding="utf-8")
        id_match = re.search(r"\*\*([F][-][A-Z]+[-][0-9]+(?:\.[0-9]+)*)\*\*", text)
        if id_match:
            form_id = id_match.group(1)
        name_match = re.search(r"# FORM_[A-Z_]+: (.+)", text)
        name = name_match.group(1).strip() if name_match else f.stem
        size_kb = f.stat().st_size // 1024
        registry_rows.append([form_id, f.name, name, str(size_kb)])
    add_table(doc, registry_rows)

    add_page_break(doc)

    # ── §3-N. Каждая форма ──
    for form_idx, form_file in enumerate(form_files, 3):
        # Заголовок формы
        text = form_file.read_text(encoding="utf-8")
        name_match = re.search(r"# FORM_[A-Z_]+: (.+)", text)
        form_name = name_match.group(1).strip() if name_match else form_file.stem
        id_match = re.search(r"\*\*([F][-][A-Z]+[-][0-9]+(?:\.[0-9]+)*)\*\*", text)
        form_id = id_match.group(1) if id_match else "F-HMP-?"

        add_heading(doc, f"{form_idx} {form_id}: {form_name}", 1)

        # Парсим и рендерим секции формы
        blocks = parse_form_md(form_file)
        for kind, payload in blocks:
            if kind == "h2":
                add_heading(doc, strip_prefix(payload), 2)
            elif kind == "h3":
                add_heading(doc, strip_prefix(payload), 3)
            elif kind == "p":
                add_body(doc, payload)
            elif kind == "ul":
                add_list_bullet(doc, payload)
            elif kind == "table":
                if len(payload) > 0:
                    # Подпись таблицы — обобщённая
                    n_rows = len(payload) - 1
                    add_table_caption_above(doc, f"Спецификация ({n_rows} строк)")
                    add_table(doc, payload)
            elif kind == "code":
                add_code_block(doc, payload)

        # Разрыв страницы между формами
        if form_idx - 2 < len(form_files):
            add_page_break(doc)

    add_page_numbers(doc)
    doc.save(out_path)
    return post_process_docx(out_path)


def main():
    project_root = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
    forms_dir = project_root / "artifacts" / "05-forms"
    out_path = forms_dir / "FORMS_SPEC.docx"

    if not forms_dir.exists():
        print(f"Not found: {forms_dir}", file=sys.stderr)
        sys.exit(1)

    ok = build_docx(forms_dir, out_path)
    if ok:
        print(f"\nOK: {out_path} ({out_path.stat().st_size // 1024} KB)")
    else:
        print("FAIL: XML-валидность нарушена", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()