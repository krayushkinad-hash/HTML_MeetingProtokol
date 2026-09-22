#!/usr/bin/env python3
"""Generate EPIC-DOCX.docx — обзор эпика DOCX-экспорта со скриншотами.

Содержит:
1. Титульный лист (User Story)
2. Сводка по эпику (3 US, 18 SP, приоритеты)
3. Детальные карточки US-015, US-016, US-017
4. Скриншоты UI (мокапы): Live Mode (3 панели), Календарь, Главная страница
5. Технические заметки + Definition of Done

Стандарт оформления — Алексей Краюшкин (TNR 12, заголовки, таблицы).
"""
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


# ─────────────────────────────────────────────────────────────────────────
# Очистка текста
# ─────────────────────────────────────────────────────────────────────────

def clean_text(s: str) -> str:
    s = s.replace("—", "-").replace("–", "-")
    s = s.replace("\xa0", " ")
    s = re.sub(r" {2,}", " ", s)
    return s.strip()


# ─────────────────────────────────────────────────────────────────────────
# Стили (Краюшкин)
# ─────────────────────────────────────────────────────────────────────────

def setup_normal_style(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    for old in rPr.findall(qn("w:lang")):
        rPr.remove(old)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")
    rPr.append(lang)
    pf = normal.paragraph_format
    pf.first_line_indent = Mm(12.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def setup_caption_style(doc):
    styles = doc.styles
    try:
        cap = styles["Подпись"]
    except KeyError:
        cap = styles.add_style("Подпись", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = styles["Normal"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(12)
    cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    pf = cap.paragraph_format
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_with_next = True


def setup_figure_caption_style(doc):
    """Стиль для подписей рисунков (моноширинный шрифт не нужен, но подпись рисунка отличается от подписи таблицы — ставится под блоком)."""
    styles = doc.styles
    try:
        cap = styles["Подпись рисунка"]
    except KeyError:
        try:
            cap = styles.add_style("Подпись рисунка", WD_STYLE_TYPE.PARAGRAPH)
            cap.base_style = styles["Normal"]
        except Exception:
            cap = styles["Подпись"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(12)
    cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    pf = cap.paragraph_format
    pf.first_line_indent = Mm(0)
    pf.space_before = Pt(2)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.keep_with_next = False


def set_page_margins(doc):
    for section in doc.sections:
        section.left_margin = Mm(20)
        section.right_margin = Mm(10)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)


def setup_headings(doc):
    sizes = {1: 16, 2: 14, 3: 12}
    for lvl, sz in sizes.items():
        try:
            s = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "Times New Roman")
        rFonts.set(qn("w:cs"), "Times New Roman")
        pf = s.paragraph_format
        pf.first_line_indent = Mm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


# ─────────────────────────────────────────────────────────────────────────
# Добавление элементов
# ─────────────────────────────────────────────────────────────────────────

def add_document_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(clean_text(text))


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_bullet(doc, items):
    for indent, txt in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(7.5 * (1 + indent // 2))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_number(doc, items):
    for txt in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_seq_field(paragraph, name="Таблица"):
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {name} \\* ARABIC "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def add_table_caption_above(doc, caption_text):
    p = doc.add_paragraph(style="Подпись")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run1 = p.add_run(f"{clean_text(caption_text)} приведено в таблице (Таблица ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Таблица")
    run2 = p.add_run(").")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_table(doc, rows):
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Light Grid Accent 1"
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean_text(cell_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return table


def add_figure_caption_below(doc, caption_text):
    """Подпись рисунка — под изображением, по центру, формат «Рисунок N. <Название>»."""
    p = doc.add_paragraph(style="Подпись рисунка")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    run1 = p.add_run("Рисунок ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Рисунок")
    run2 = p.add_run(f". {clean_text(caption_text)}")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_picture(doc, image_path: Path, caption_text: str, max_width_cm: float = 15.0):
    """Встраивает картинку + подпись рисунка под ней."""
    if not image_path.exists():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(max_width_cm))
    add_figure_caption_below(doc, caption_text)
    return True


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


# ─────────────────────────────────────────────────────────────────────────
# Пост-обработка DOCX
# ─────────────────────────────────────────────────────────────────────────

def post_process_docx(docx_path: Path) -> bool:
    tmp = docx_path.with_suffix(".tmp.docx")
    seen_n = False
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/numbering.xml":
                    if seen_n:
                        continue
                    seen_n = True
                zout.writestr(item.filename, zin.read(item.filename))
    shutil.move(tmp, docx_path)
    with zipfile.ZipFile(docx_path, "r") as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(name))
                except ET.ParseError as e:
                    print(f"INVALID: {name}: {e}")
                    return False
    return True


# ─────────────────────────────────────────────────────────────────────────
# Сборка документа
# ─────────────────────────────────────────────────────────────────────────

def build_docx(out_path: Path, project_root: Path):
    screenshots_dir = project_root / "assets" / "screenshots"
    doc = Document()
    setup_normal_style(doc)
    setup_headings(doc)
    setup_caption_style(doc)
    setup_figure_caption_style(doc)
    set_page_margins(doc)

    doc.core_properties.author = "Алексей Краюшкин"
    doc.core_properties.title = "EPIC-DOCX HTML_MeetingProtokol"

    # ── Титульный лист ──
    add_document_title(doc, "EPIC-DOCX: Экспорт протокола в DOCX")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("Проект HTML_MeetingProtokol")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("Версия 1.0 · 2026-09-14")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_page_break(doc)

    # ── §1. Метаинформация эпика ──
    add_heading(doc, "1 Метаинформация эпика", 1)
    add_table_caption_above(doc, "Метаинформация эпика EPIC-DOCX")
    add_table(doc, [
        ["Поле", "Значение"],
        ["Эпик", "EPIC-DOCX"],
        ["Проект", "HTML_MeetingProtokol"],
        ["Владелец", "Алексей Краюшкин"],
        ["Story Points (сумма)", "18"],
        ["User Stories", "3"],
        ["Дата", "2026-09-14"],
    ])

    # ── §2. Цель эпика ──
    add_heading(doc, "2 Цель эпика", 1)
    add_body(doc,
        "Превратить распознанный протокол встречи в готовый к отправке или печати "
        "документ формата DOCX, который: оформлен по стандарту Краюшкина (TNR 12, "
        "заголовки 1-2-3, таблицы в стиле Сетка таблицы); содержит таймкоды реплик "
        "с гиперссылками на локальное видео; реплики сгруппированы по ораторам; "
        "опционально содержит встроенные скриншоты экрана демонстранта; генерируется "
        "за короткое время (не более 30 секунд для протокола без скриншотов)."
    )

    # ── §3. User Stories ──
    add_heading(doc, "3 User Stories эпика", 1)
    add_table_caption_above(doc, "Сводная таблица US эпика EPIC-DOCX")
    add_table(doc, [
        ["ID", "US-текст", "Приоритет", "SP", "Персона"],
        ["US-015",
         "Как Алексей (P-01), я хочу экспортировать протокол в DOCX с таймкодами и репликами по ораторам, чтобы отправить участникам встречи",
         "Must", "5", "P-01, P-03"],
        ["US-016",
         "Как Алексей (P-01), я хочу, чтобы в DOCX были встроены скриншоты экрана демонстранта с привязкой к моменту реплики",
         "Should", "8", "P-01"],
        ["US-017",
         "Как Алексей (P-01), я хочу, чтобы DOCX содержал гиперссылку открыть видео с HH:MM:SS на локальный файл, чтобы участник мог перейти к нужному моменту",
         "Should", "5", "P-01"],
    ])

    # ── §4. Контекст в бизнес-процессах ──
    add_heading(doc, "4 Контекст в бизнес-процессах", 1)
    add_heading(doc, "4.1 As-Is (текущее состояние)", 2)
    add_body(doc,
        "Алексей вручную открывает Word или Google Docs, копирует текст из блокнота "
        "или из прослушки записи встречи. Расставляет таймкоды вручную, часто с "
        "ошибками. Сохраняет документ, отправляет участникам по электронной почте. "
        "Время на подготовку одного протокола: 30-60 минут для 1-часовой встречи. "
        "Проблемы: пропускает реплики, путает ораторов, нет гиперссылок на видео, "
        "форматирование часто отличается от стиля других документов."
    )
    add_heading(doc, "4.2 To-Be (целевое состояние)", 2)
    add_body(doc,
        "Алексей нажимает кнопку Экспорт - DOCX в интерфейсе HTML_MeetingProtokol. "
        "Backend формирует DOCX по стандарту Краюшкина за не более чем 30 секунд. "
        "Документ содержит все необходимые элементы: таймкоды, имена ораторов, "
        "гиперссылки на локальное видео, опционально скриншоты экрана. Файл "
        "отправляется участникам встречи. Время на подготовку одного протокола: "
        "30 секунд плюс 1 минута на отправку. Выигрыш по сравнению с ручным "
        "режимом: с 30-60 минут до 1.5 минут, что в 20-40 раз быстрее."
    )

    # ── §5. Скриншоты интерфейса ──
    add_heading(doc, "5 Скриншоты интерфейса", 1)
    add_body(doc,
        "Ниже представлены макеты ключевых экранов HTML_MeetingProtokol, "
        "которые связаны с эпиком EPIC-DOCX (результат работы этих экранов "
        "экспортируется в DOCX)."
    )

    add_picture(doc, screenshots_dir / "live_mode.png",
                "Режим Live Mode: три панели (видео участников, темы встречи, живая транскрипция) — данные из этого режима экспортируются в DOCX по US-015, US-016, US-017",
                max_width_cm=15.5)

    add_picture(doc, screenshots_dir / "calendar.png",
                "Календарь прошлых встреч: список протоколов на выбранную дату становится источником для экспорта",
                max_width_cm=14.5)

    add_picture(doc, screenshots_dir / "upload.png",
                "Главная страница: список всех протоколов с быстрым доступом к экспорту DOCX",
                max_width_cm=14.5)

    add_page_break(doc)

    # ── §6. Детальная карточка US-015 ──
    add_heading(doc, "6 US-015: Экспорт протокола в DOCX с таймкодами", 1)
    add_table_caption_above(doc, "Метаданные US-015")
    add_table(doc, [
        ["Поле", "Значение"],
        ["ID", "US-015"],
        ["Эпик", "EPIC-DOCX"],
        ["Приоритет", "Must"],
        ["Story Points", "5"],
        ["Персона", "P-01 Алексей (Windows), P-03 Гость"],
        ["Форма", "ExportDialog"],
        ["API", "POST /api/v1/export/docx"],
    ])

    add_heading(doc, "6.1 User Story", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    run = p.add_run("Как ")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run("Алексей (P-01)")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run(", я хочу экспортировать протокол в DOCX с таймкодами и репликами по ораторам, чтобы отправить участникам встречи.")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    add_heading(doc, "6.2 Acceptance Criteria", 2)
    add_list_number(doc, [
        "AC-1: открывается диалог с настройками: включить таймкоды, скриншоты, гиперссылки, группировать по ораторам.",
        "AC-2: формируется DOCX-файл по стандарту Краюшкина (TNR 12, заголовки, таблицы Сетка таблицы).",
        "AC-3: в DOCX есть титульный лист с названием и датой встречи, список участников, повестка, реплики с таймкодами.",
        "AC-4: рядом с таймкодом стоит гиперссылка для открытия видео с этой секунды.",
        "AC-5: при наличии скриншотов они встроены в DOCX с привязкой к моменту реплики.",
    ])

    add_page_break(doc)

    # ── §7. US-016 ──
    add_heading(doc, "7 US-016: Скриншоты экрана в DOCX", 1)
    add_table_caption_above(doc, "Метаданные US-016")
    add_table(doc, [
        ["Поле", "Значение"],
        ["ID", "US-016"],
        ["Эпик", "EPIC-DOCX"],
        ["Приоритет", "Should"],
        ["Story Points", "8"],
        ["Персона", "P-01 Алексей (Windows)"],
        ["Форма", "ExportDialog"],
        ["API", "POST /api/v1/export/docx?screenshots=true"],
    ])

    add_heading(doc, "7.1 User Story", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    run = p.add_run("Как ")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run("Алексей (P-01)")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run(", я хочу, чтобы в DOCX были встроены скриншоты экрана демонстранта с привязкой к моменту реплики, чтобы участники видели, что показывалось в момент обсуждения.")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    add_heading(doc, "7.2 Acceptance Criteria", 2)
    add_list_number(doc, [
        "AC-1: при включенной опции скриншоты встроены в DOCX после реплики, во время которой был сделан скриншот.",
        "AC-2: под изображением стоит подпись Рисунок N. Скриншот экрана в момент реплики.",
        "AC-3: размер скриншота масштабируется до 15 см по ширине с сохранением пропорций.",
        "AC-4: итоговый размер DOCX не превышает 50 МБ (иначе предупреждение пользователю).",
        "AC-5: при отсутствии скриншотов опция автоматически отключается без ошибок.",
    ])

    add_page_break(doc)

    # ── §8. US-017 ──
    add_heading(doc, "8 US-017: Гиперссылки на видео в DOCX", 1)
    add_table_caption_above(doc, "Метаданные US-017")
    add_table(doc, [
        ["Поле", "Значение"],
        ["ID", "US-017"],
        ["Эпик", "EPIC-DOCX"],
        ["Приоритет", "Should"],
        ["Story Points", "5"],
        ["Персона", "P-01 Алексей (Windows)"],
        ["Форма", "ExportDialog"],
        ["API", "POST /api/v1/export/docx"],
    ])

    add_heading(doc, "8.1 User Story", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    run = p.add_run("Как ")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run("Алексей (P-01)")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    run = p.add_run(", я хочу, чтобы в DOCX рядом с каждой репликой стояла гиперссылка вида HH:MM:SS на локальный видеофайл, чтобы участник мог кликнуть и сразу перейти к нужному моменту разговора.")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    add_heading(doc, "8.2 Acceptance Criteria", 2)
    add_list_number(doc, [
        "AC-1: рядом с репликой стоит кликабельный текст вида HH:MM:SS (синий, подчеркнутый).",
        "AC-2: клик открывает системный видеоплеер по умолчанию с воспроизведением с указанной секунды.",
        "AC-3: гиперссылка указывает на относительный путь к видеофайлу плюс параметр с фрагментом времени.",
        "AC-4: при отсутствии видеофайла клик показывает понятное сообщение об ошибке.",
        "AC-5: в сопроводительном письме явно указано, что для работы гиперссылок нужно приложить видеофайл.",
    ])

    add_page_break(doc)

    # ── §9. Технические заметки ──
    add_heading(doc, "9 Технические заметки", 1)
    add_body(doc,
        "Стек реализации: backend на Python с библиотекой python-docx версии 1.1.0 и выше. "
        "Шаблон оформления - скрипт scripts/render_vision_docx.py, который уже есть в "
        "проекте для генерации VISION.docx. Этот скрипт адаптируется для протокола: "
        "добавляются функции add_hyperlink (обертка для гиперссылок через OXML), "
        "add_picture (встраивание скриншотов с автоматическим подгоном размера), "
        "add_figure_caption_below (подпись рисунка под блоком)."
    )
    add_body(doc,
        "Размер файла: без скриншотов DOCX занимает 100-500 КБ, со скриншотами "
        "(10-30 штук) - от 5 до 50 МБ. Гиперссылки формируются по шаблону "
        "file:./video.mp4 с фрагментом времени. В сопроводительном письме "
        "указывается, что для работы гиперссылок нужно приложить видеофайл к письму."
    )
    add_body(doc,
        "Совместимость со стандартом Краюшкина: используется готовый шаблон "
        "scripts/render_vision_docx.py как основа. Шрифт Times New Roman 12 пт, "
        "поля 20/10/10/10 мм, заголовки документа (18 пт жирный по центру), "
        "Heading 1 (16 пт), Heading 2 (14 пт), Heading 3 (12 пт), таблицы в стиле "
        "Сетка таблицы с повтором заголовков, нумерация страниц внизу по центру."
    )

    # ── §10. Definition of Done ──
    add_heading(doc, "10 Definition of Done", 1)
    add_body(doc, "Эпик считается завершенным, когда выполнены все перечисленные ниже пункты.")
    add_list_number(doc, [
        "US-015, US-016, US-017 реализованы и проходят все Acceptance Criteria.",
        "Сформированный DOCX проходит все 33 пункта чек-листа приемки по стандарту Краюшкина.",
        "Файл корректно открывается в Word 2019+, LibreOffice 7+, Google Docs.",
        "XML-валидность всех файлов внутри DOCX-архива проверена через xml.etree.ElementTree.fromstring.",
        "Размер файла не превышает 10 МБ для протокола без скриншотов и 50 МБ со скриншотами.",
        "Покрытие unit-тестами не менее 70 процентов для скрипта генерации.",
        "Проверена работа на Astra Linux: шрифт Times New Roman установлен, документ открывается в LibreOffice.",
    ])

    add_page_numbers(doc)
    doc.save(out_path)
    ok = post_process_docx(out_path)
    return ok


def main():
    project_root = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
    out_path = project_root / "artifacts" / "03-user-stories" / "EPIC-DOCX.docx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ok = build_docx(out_path, project_root)
    if ok:
        print(f"OK: {out_path} ({out_path.stat().st_size // 1024} KB)")
    else:
        print("FAIL: XML-валидность нарушена", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()