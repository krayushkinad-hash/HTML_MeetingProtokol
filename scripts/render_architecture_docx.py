#!/usr/bin/env python3
"""Render ARCHITECTURE.md → DOCX по Краюшкину + генерация диаграмм PNG."""
import re
import shutil
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from PIL import Image
from playwright.sync_api import sync_playwright


# ─── Helpers ───

def clean_text(s):
    s = s.replace("—", "-").replace("–", "-")
    s = s.replace("\xa0", " ")
    s = re.sub(r" {2,}", " ", s)
    return s.strip()


def setup_krayushkin_styles(doc):
    """Стили по стандарту Краюшкина (TNR, поля, заголовки, подписи)."""
    from docx.enum.style import WD_STYLE_TYPE

    # Create custom styles if not exist
    styles = doc.styles
    if "Подпись" not in [s.name for s in styles]:
        cap = styles.add_style("Подпись", WD_STYLE_TYPE.PARAGRAPH)
        cap.base_style = styles["Normal"]
        cap.font.name = "Times New Roman"
        cap.font.size = Pt(12)
        cap.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        cap.paragraph_format.first_line_indent = Mm(0)
        cap.paragraph_format.space_before = Pt(12)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    for old in rPr.findall(qn("w:lang")):
        rPr.remove(old)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")
    lang.set(qn("w:bidi"), "ru-RU")
    rPr.append(lang)
    pf = normal.paragraph_format
    pf.first_line_indent = Mm(12.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    for lvl, sz in [(1, 16), (2, 14), (3, 12)]:
        try:
            s = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        s.font.name = "Times New Roman"
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        rPr = s.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "Times New Roman")
        rFonts.set(qn("w:cs"), "Times New Roman")
        pf = s.paragraph_format
        pf.first_line_indent = Mm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.15
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def set_page_margins(doc):
    for section in doc.sections:
        section.left_margin = Mm(20)
        section.right_margin = Mm(10)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)


def add_page_numbers(doc):
    for section in doc.sections:
        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def add_document_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(clean_text(text))


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(clean_text(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_list_bullet(doc, items):
    for indent, txt in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(7.5 * (1 + indent // 2))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(clean_text(txt))
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_seq_field(paragraph, name):
    run = paragraph.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {name} \\* ARABIC "
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def add_table_caption_above(doc, caption_text):
    p = doc.add_paragraph(style="Подпись")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run1 = p.add_run(f"{clean_text(caption_text)} приведено в таблице (Таблица ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Таблица")
    run2 = p.add_run(").")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_figure_caption_below(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    run1 = p.add_run("Рисунок ")
    run1.font.name = "Times New Roman"; run1.font.size = Pt(12); run1.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    add_seq_field(p, "Рисунок")
    run2 = p.add_run(f". {clean_text(caption_text)}")
    run2.font.name = "Times New Roman"; run2.font.size = Pt(12); run2.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_picture(doc, image_path, caption_text, max_width_cm=15.5):
    if not image_path.exists():
        return False
    with Image.open(image_path) as img:
        w_px, h_px = img.size
    target_w = max_width_cm
    target_h = target_w * (h_px / w_px)
    if target_h > 20:
        target_h = 20
        target_w = target_h * (w_px / h_px)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(target_w), height=Cm(target_h))
    add_figure_caption_below(doc, caption_text)
    return True


def add_table(doc, rows):
    if not rows:
        return None
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Mm(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(clean_text(cell_text))
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


# ─── Mermaid rendering ───

CHROMIUM = "/root/.cache/ms-playwright/chromium-1228/chrome-linux64/chrome"

HTML_TEMPLATE = """<!DOCTYPE html>
<html><head><meta charset="UTF-8">
<style>body{{margin:20px;background:white;font-family:sans-serif;}}.mermaid{{font-size:14px;}}</style>
<script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
</head><body>
<pre class="mermaid">
{code}
</pre>
<script>
mermaid.initialize({{ startOnLoad: false, theme: 'default', er: {{ htmlLabels: true }} }});
mermaid.run().then(() => {{ window.rendered = true; }});
</script>
</body></html>"""


def render_mermaid(code, output_path, viewport_w=2400, viewport_h=1800):
    tmp_html = output_path.parent / "_temp_arch.html"
    tmp_html.write_text(HTML_TEMPLATE.format(code=code), encoding="utf-8")
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROMIUM)
        page = browser.new_page(viewport={"width": viewport_w, "height": viewport_h})
        try:
            page.goto(f"file://{tmp_html.resolve()}", wait_until="domcontentloaded")
            try:
                page.wait_for_function("window.rendered === true", timeout=15000)
            except Exception:
                page.wait_for_selector(".mermaid svg", timeout=15000)
            page.wait_for_timeout(1000)
            page.locator(".mermaid").screenshot(path=str(output_path))
        finally:
            page.close()
            browser.close()
    tmp_html.unlink()


# ─── MD parser ───

def parse_md(path):
    text = path.read_text(encoding="utf-8")
    blocks = []
    lines = text.split("\n")
    i = 0
    in_mermaid = False
    in_code = False
    code_lines = []
    mermaid_code = []
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i].rstrip()

        # Mermaid → в PNG
        if line.strip().startswith("```mermaid"):
            in_mermaid = True
            mermaid_code = []
            i += 1
            continue
        if in_mermaid:
            if line.strip() == "```":
                in_mermaid = False
                blocks.append(("mermaid", "\n".join(mermaid_code)))
            else:
                mermaid_code.append(line)
            i += 1
            continue

        # SQL/code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                blocks.append(("code", "\n".join(code_lines)))
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # H1 — пропускаем (свой заголовок)
        if re.match(r"^#\s+", line) and not re.match(r"^##\s+", line):
            i += 1
            continue

        # H2 / H3
        m = re.match(r"^(#{2,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            blocks.append((f"h{level}", title))
            i += 1
            continue

        # Таблица
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|", lines[i + 1].strip()):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            rows = []
            for row in tbl:
                if re.match(r"^\|[\s\-:|]+\|$", row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                rows.append(cells)
            blocks.append(("table", rows))
            continue

        # Буллет
        if re.match(r"^\s*-\s+", line):
            items = []
            while i < len(lines) and re.match(r"^\s*-\s+", lines[i]):
                raw = lines[i]
                indent = len(raw) - len(raw.lstrip())
                txt = re.sub(r"^\s*-\s+", "", raw).rstrip()
                items.append((indent, txt))
                i += 1
            blocks.append(("ul", items))
            continue

        # Параграф
        if line.strip():
            blocks.append(("p", line))
        i += 1

    return blocks


def strip_prefix(text):
    return re.sub(r"^\d+(\.\d+)*\.?\s+", "", text).strip()


# ─── Post-processing ───

def post_process_docx(docx_path):
    tmp = docx_path.with_suffix(".tmp.docx")
    seen_n = False
    with zipfile.ZipFile(docx_path, "r") as zin:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename == "word/numbering.xml":
                    if seen_n:
                        continue
                    seen_n = True
                zout.writestr(item.filename, zin.read(item.filename))
    shutil.move(tmp, docx_path)
    with zipfile.ZipFile(docx_path, "r") as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    ET.fromstring(z.read(name))
                except ET.ParseError as e:
                    print(f"INVALID: {name}: {e}")
                    return False
    return True


# ─── Main ───

def main():
    project_root = Path(sys.argv[1] if len(sys.argv) > 1 else ".").resolve()
    md_path = project_root / "artifacts" / "11-architecture" / "ARCHITECTURE.md"
    diagrams_dir = project_root / "artifacts" / "11-architecture" / "diagrams"
    diagrams_dir.mkdir(parents=True, exist_ok=True)
    out_path = md_path.parent / "ARCHITECTURE.docx"

    if not md_path.exists():
        print(f"NotFound: {md_path}", file=sys.stderr)
        sys.exit(1)

    print("Шаг 1 Парсинг ARCHITECTURE.md...")
    blocks = parse_md(md_path)

    # Подсчёт Mermaid-блоков
    mermaid_count = sum(1 for kind, _ in blocks if kind == "mermaid")
    print(f"  Найдено {len(blocks)} блоков, {mermaid_count} Mermaid-диаграмм")

    print("\nШаг 2 Генерация диаграмм PNG...")
    doc = Document()
    setup_krayushkin_styles(doc)
    set_page_margins(doc)

    doc.core_properties.author = "Алексей Краюшкин"
    doc.core_properties.title = "Architecture HTML_MeetingProtokol"

    add_document_title(doc, "Architecture HTML_MeetingProtokol")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("Документ архитектуры приложения · версия 1.0")
    run.font.name = "Times New Roman"; run.font.size = Pt(14); run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    run = p.add_run("2026-09-14")
    run.font.name = "Times New Roman"; run.font.size = Pt(12); run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    add_page_break(doc)

    # Render & embed Mermaid
    diag_idx = 0
    table_counter = 0
    for kind, payload in blocks:
        if kind == "h2":
            add_heading(doc, strip_prefix(payload), 1)
        elif kind == "h3":
            add_heading(doc, strip_prefix(payload), 2)
        elif kind == "p":
            add_body(doc, payload)
        elif kind == "ul":
            add_list_bullet(doc, payload)
        elif kind == "table":
            if len(payload) > 0:
                table_counter += 1
                add_table_caption_above(doc, f"Данные раздела")
                add_table(doc, payload)
        elif kind == "code":
            for line in payload.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Mm(0)
                p.paragraph_format.left_indent = Mm(12.5)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                run = p.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        elif kind == "mermaid":
            diag_idx += 1
            png_path = diagrams_dir / f"diagram_{diag_idx}.png"
            print(f"  Диаграмма {diag_idx}...")
            try:
                render_mermaid(payload, png_path, viewport_w=2400, viewport_h=1800)
                # Find context (look at previous block)
                context_title = "Диаграмма"
                if diag_idx >= 1 and blocks.index(("mermaid", payload)) > 0:
                    # Get heading before
                    idx = blocks.index(("mermaid", payload))
                    for j in range(idx - 1, -1, -1):
                        if blocks[j][0] in ("h2", "h3"):
                            context_title = strip_prefix(blocks[j][1])
                            break
                add_picture(doc, png_path, f"{context_title} (диаграмма {diag_idx})", max_width_cm=16.0)
                print(f"    → {png_path.name} ({png_path.stat().st_size // 1024} KB)")
            except Exception as e:
                print(f"    ERROR: {e}")
                # Add placeholder text
                add_body(doc, f"[Не удалось отрендерить диаграмму {diag_idx}: {e}]")

    add_page_numbers(doc)
    doc.save(out_path)
    print(f"\nПост-обработка DOCX...")
    ok = post_process_docx(out_path)
    if ok:
        print(f"\nOK: {out_path} ({out_path.stat().st_size // 1024} KB)")
    else:
        print("FAIL: XML-валидность нарушена", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()